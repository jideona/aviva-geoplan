"""Parse the Aviva Wuye route-analysis workbook (field survey, Feb 2025).

Pure: bytes in, records out. Two sheets matter — recorded street names with
measured lengths, and estate unit counts, which are the only calibration input
the premises model has.
"""
import io
import re
from dataclasses import dataclass

from openpyxl import load_workbook

from app.domain.premises_model import band_for_units


class WorkbookParseError(ValueError):
    """Message is safe to show the user."""


@dataclass
class RecordedStreetRow:
    name: str
    length_m: float | None


@dataclass
class EstateRow:
    estate_name: str
    building_count: int
    units_per_building: int
    typology: str
    street_hint: str | None


# "423 Magnus Abe St", "Plot 386 Jaja Nwachukwu St" -> the street it sits on.
_STREET_HINT = re.compile(
    r"(?:plot\s+)?\d+[a-z]?\s+(.+?(?:st|street|cres|crescent|road|rd|way|blvd|"
    r"boulevard|avenue|ave|close|cl))\b", re.I)


def _street_hint(label: str) -> str | None:
    m = _STREET_HINT.search(label.strip())
    return m.group(1).strip() if m else None


def parse(data: bytes) -> tuple[list[RecordedStreetRow], list[EstateRow]]:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise WorkbookParseError(f"Could not read the workbook: {exc}") from exc

    streets = _streets(wb)
    estates = _estates(wb)
    if not streets and not estates:
        raise WorkbookParseError(
            "No 'Street Count' or 'Estate Count' sheet found. This importer "
            "expects the Aviva route-analysis workbook layout."
        )
    return streets, estates


def _sheet(wb, wanted: str):
    for name in wb.sheetnames:
        if name.strip().lower() == wanted:
            return wb[name]
    return None


def _streets(wb) -> list[RecordedStreetRow]:
    ws = _sheet(wb, "street count")
    if ws is None:
        return []
    seen: set[str] = set()
    out: list[RecordedStreetRow] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        name = str(row[0]).strip()
        if not name or name.lower() == "street":
            continue
        length = float(row[1]) if len(row) > 1 and isinstance(
            row[1], (int, float)) else None
        key = name.lower()
        if key in seen:
            # The workbook repeats a street; keep the first measurement.
            continue
        seen.add(key)
        out.append(RecordedStreetRow(name=name, length_m=length))
    return out


def _estates(wb) -> list[EstateRow]:
    ws = _sheet(wb, "estate count")
    if ws is None:
        return []
    out: list[EstateRow] = []
    current: str | None = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        # Column B carries the estate name; blank means "same estate as above",
        # which is how the surveyor recorded mixed building types on one site.
        if len(row) > 1 and row[1]:
            current = str(row[1]).strip()
        if current is None:
            continue
        buildings = row[2] if len(row) > 2 else None
        units = row[3] if len(row) > 3 else None
        if not isinstance(buildings, (int, float)) or not isinstance(
                units, (int, float)):
            continue
        b, u = int(buildings), int(units)
        if b <= 0 or u <= 0:
            continue
        out.append(EstateRow(
            estate_name=current, building_count=b, units_per_building=u,
            typology=band_for_units(u).value, street_hint=_street_hint(current),
        ))
    return out
