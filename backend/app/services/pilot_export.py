"""Buildable pilot design pack — Excel BOQ and schedule."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY="FF0D1B4B"; BRAND="FF1A6FA8"; TEAL="FF00C9A7"; AMBER="FFB45309"
STEEL="FF5A739A"; LIGHT="FFF7FAFD"


def _hdr(ws, headers, widths, row=1, fill=NAVY):
    for i,(h,w) in enumerate(zip(headers,widths),start=1):
        c=ws.cell(row=row,column=i,value=h)
        c.font=Font(name="Space Mono",size=9,bold=True,color="FFFFFFFF")
        c.fill=PatternFill("solid",fgColor=fill)
        c.alignment=Alignment(vertical="center",wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[row].height=24


def build_workbook(plan: dict, project_name: str) -> bytes:
    wb=Workbook()
    ws=wb.active; ws.title="Pilot summary"
    ws.column_dimensions["A"].width=34; ws.column_dimensions["B"].width=34
    ws["A1"]="AVIVA NETWORX"; ws["A1"].font=Font(name="DM Sans",size=16,bold=True,color=NAVY)
    ws["A2"]=f"{project_name} — pilot design & bill of quantities"
    ws["A2"].font=Font(name="DM Sans",size=12,color=BRAND)

    rows=[
        ("OLT", f'{plan["olt_pon_ports"]}-port, {plan["equipment"]["olt"]["ports_used"]} ports used'),
        ("",""),
        ("Phase 1 — connectorised core", ""),
        ("  FATs (pre-terminated)", plan["equipment"]["fats"]["connectorised"]),
        ("  premises", plan["phase1_premises"]),
        ("  build", "plug-and-play, no field splicing"),
        ("",""),
        ("Phase 2 — conventional extension", ""),
        ("  FATs (spliced)", plan["equipment"]["fats"]["conventional"]),
        ("  premises", plan["phase2_premises"]),
        ("  build", "two-stage 1:32, spliced"),
        ("",""),
        ("TOTAL PILOT PREMISES", plan["total_premises"]),
        ("FATs total", plan["equipment"]["fats"]["total"]),
        ("FATs beyond pilot scope", plan["unserved_fats"]),
    ]
    r=4
    for label,value in rows:
        if label:
            bold = label.startswith("Phase") or label.startswith("TOTAL")
            ws.cell(row=r,column=1,value=label).font=Font(
                name="DM Sans",size=10,bold=bold,color=NAVY if bold else STEEL)
            ws.cell(row=r,column=2,value=value).font=Font(name="DM Sans",size=10,color=NAVY)
        r+=1
    if plan["warnings"]:
        r+=1
        for w in plan["warnings"]:
            c=ws.cell(row=r,column=1,value="⚠ "+w)
            c.font=Font(name="DM Sans",size=9,italic=True,color=AMBER)
            c.alignment=Alignment(wrap_text=True)
            ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=2)
            ws.row_dimensions[r].height=28; r+=1

    # --- BOQ: splitters + procurement ---
    bs=wb.create_sheet("Splitter BOQ")
    _hdr(bs,["Item","Role","Required","With spares","In stock","BUY"],
         [16,32,10,12,10,8])
    ri=2
    for l in plan["equipment"]["splitters"]:
        vals=[l["item"],l["role"],l["required"],l["with_spares"],l["in_stock"],l["buy"]]
        for j,v in enumerate(vals,start=1):
            c=bs.cell(row=ri,column=j,value=v); c.font=Font(name="DM Sans",size=9)
            if j==6 and l["buy"]>0: c.font=Font(name="Space Mono",size=9,bold=True,color=AMBER)
        ri+=1
    bs.cell(row=ri+1,column=1,value="Only the 1:4 primaries need procurement — a "
            "trivial spend that keeps the optical budget comfortable and fills "
            "the OLT.").font=Font(name="DM Sans",size=9,italic=True,color=STEEL)
    bs.merge_cells(start_row=ri+1,start_column=1,end_row=ri+1,end_column=6)

    # --- connectorised cable usage ---
    cs=wb.create_sheet("Connectorised cables")
    _hdr(cs,["Cable assembly","Used"],[28,10],fill=TEAL)
    ri=2
    for cable,n in sorted(plan["equipment"]["connectorised_cables_used"].items()):
        cs.cell(row=ri,column=1,value=cable).font=Font(name="DM Sans",size=9)
        cs.cell(row=ri,column=2,value=n).font=Font(name="DM Sans",size=9)
        ri+=1

    # --- FAT schedule ---
    fs=wb.create_sheet("FAT schedule")
    _hdr(fs,["FAT","Phase","Type","Premises","Route m","Cable / build"],
         [14,7,15,10,10,24])
    fs.freeze_panes="A2"
    allf=sorted(plan["phase1_fats"]+plan["phase2_fats"],
                key=lambda f:(f["phase"],f["fat_code"]))
    for i,f in enumerate(allf,start=2):
        vals=[f["fat_code"],f["phase"],f["fat_type"],f["premises"],f["route_m"],
              f["cable"] or "spliced FAT + 1:8"]
        for j,v in enumerate(vals,start=1):
            c=fs.cell(row=i,column=j,value=v); c.font=Font(name="DM Sans",size=9)
            if i%2==0: c.fill=PatternFill("solid",fgColor=LIGHT)
    fs.auto_filter.ref=f"A1:F{len(allf)+1}"

    out=io.BytesIO(); wb.save(out); return out.getvalue()
