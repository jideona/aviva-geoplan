"""Register exports with attribution and licence gating.

SRD FR-EXP-001 to FR-EXP-004, FR-IMP-032, FR-IMP-033.
"""
import csv
import io
import json
import uuid
from datetime import datetime, timezone

from geoalchemy2.shape import to_shape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from shapely.geometry import mapping
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db.models.building import Building
from app.db.models.project import Project
from app.db.models.provenance import DataSource
from app.db.models.street import Street
from app.db.models.user import User
from app.domain.attribution import ExportPurpose, Manifest, SourceEntry, build
from app.services import audit_service, register_service, walkthrough_service
from app.services.register_service import RegisterFilter

# Aviva Networx Brand Identity Guide v1.0
NAVY = "FF0D1B4B"
BRAND = "FF1A6FA8"
LIGHT = "FFE8EEF6"
STEEL = "FF5A739A"

COLUMNS = [
    ("Building ID", "building_code", 22),
    ("Street code", "street_code", 14),
    ("Street name", "street_name", 26),
    ("Footprint m2", "footprint_area_sqm", 13),
    ("Perimeter m", "perimeter_m", 12),
    ("Building type", "building_type", 18),
    ("Use", "use_type", 13),
    ("Floors", "floors_reported", 8),
    ("Units surveyed", "units_surveyed", 14),
    ("Premises estimated", "premises_estimated", 17),
    ("Longitude", "lon", 12),
    ("Latitude", "lat", 12),
    ("Assignment confidence", "street_assignment_confidence", 20),
    ("Assignment distance m", "street_assignment_distance_m", 20),
    ("Source dataset", "source_dataset", 22),
    ("Licence class", "licence_class", 20),
    ("Verification", "verification_state", 16),
    ("Survey status", "survey_status", 15),
]


class ExportBlocked(PermissionError):
    """Raised where the export purpose is incompatible with source licences."""


def _rows(db: Session, project_id: uuid.UUID, f: RegisterFilter) -> list[dict]:
    out = []
    for b, street_name, street_code in register_service.query(
            db, project_id, f, limit=None):
        pt = to_shape(b.centroid)
        out.append({
            "building_code": b.building_code or "",
            "street_code": street_code or "",
            "street_name": street_name or "",
            "footprint_area_sqm": float(b.footprint_area_sqm),
            "perimeter_m": float(b.perimeter_m),
            "building_type": b.building_type,
            "use_type": b.use_type,
            "floors_reported": b.floors_reported,
            "units_surveyed": b.units_surveyed,
            "premises_estimated": b.premises_estimated,
            "lon": round(pt.x, 7),
            "lat": round(pt.y, 7),
            "street_assignment_confidence": (
                float(b.street_assignment_confidence)
                if b.street_assignment_confidence is not None else None),
            "street_assignment_distance_m": (
                float(b.street_assignment_distance_m)
                if b.street_assignment_distance_m is not None else None),
            "source_dataset": b.source_dataset or "",
            "licence_class": b.licence_class,
            "verification_state": b.verification_state,
            "survey_status": b.survey_status,
            "_geom": b.geom,
        })
    return out


def manifest(db: Session, project: Project, purpose: ExportPurpose) -> Manifest:
    rows = db.execute(
        select(DataSource.name, DataSource.licence, DataSource.licence_class,
               func.count(Building.id))
        .join(Building, Building.data_source_id == DataSource.id)
        .where(Building.project_id == project.id)
        .group_by(DataSource.name, DataSource.licence, DataSource.licence_class)
    ).all()
    sources = [SourceEntry(n, lic, lc, c) for n, lic, lc, c in rows]

    # Street name sources contribute obligations of their own.
    for src, count_ in db.execute(
        select(Street.name_source, func.count())
        .where(Street.project_id == project.id, Street.name_source.isnot(None))
        .group_by(Street.name_source)
    ).all():
        from app.domain.naming import rules_for
        try:
            r = rules_for(src)
            sources.append(SourceEntry(f"Street names — {src.replace('_', ' ')}",
                                       None, r.licence_class.value, count_))
        except (KeyError, ValueError):
            continue

    return build(project.name,
                 datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
                 purpose, sources)


def _guard(m: Manifest) -> None:
    if m.blocked:
        raise ExportBlocked(m.reason or "Export blocked by licence obligations.")


def to_csv(db: Session, user: User, project: Project, f: RegisterFilter,
           purpose: ExportPurpose) -> str:
    m = manifest(db, project, purpose)
    _guard(m)
    rows = _rows(db, project.id, f)
    buf = io.StringIO()
    for line in m.as_text().splitlines():
        buf.write(f"# {line}\n")
    buf.write(f"# Filter: {f.describe()}\n")
    w = csv.writer(buf)
    w.writerow([label for label, _, _ in COLUMNS])
    for r in rows:
        w.writerow([r[key] for _, key, _ in COLUMNS])
    _audit(db, user, project, "csv", len(rows), f, purpose)
    return buf.getvalue()


def to_xlsx(db: Session, user: User, project: Project, f: RegisterFilter,
            purpose: ExportPurpose) -> bytes:
    m = manifest(db, project, purpose)
    _guard(m)
    rows = _rows(db, project.id, f)
    summary = register_service.summary(db, project.id)

    wb = Workbook()
    _summary_sheet(wb.active, project, f, m, summary, len(rows))
    _register_sheet(wb.create_sheet("Building register"), rows)
    _sources_sheet(wb.create_sheet("Sources & attribution"), m)

    out = io.BytesIO()
    wb.save(out)
    _audit(db, user, project, "xlsx", len(rows), f, purpose)
    return out.getvalue()


def _title(ws, text: str, row: int) -> None:
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name="DM Sans", size=11, bold=True, color="FFFFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)


def _summary_sheet(ws, project: Project, f: RegisterFilter, m: Manifest,
                   s: dict, exported: int) -> None:
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 30

    ws["A1"] = "AVIVA NETWORX"
    ws["A1"].font = Font(name="DM Sans", size=16, bold=True, color=NAVY)
    ws["A2"] = "GeoPlan — Building Register"
    ws["A2"].font = Font(name="DM Sans", size=12, color=BRAND)

    rows = [
        ("Project", project.name),
        ("District", project.district),
        ("Storage CRS", "EPSG:4326"),
        ("Measurement CRS", f"EPSG:{project.metric_crs_epsg}"),
        ("Generated", m.generated),
        ("Export purpose", m.purpose.value),
        ("Filter applied", f.describe()),
        ("Rows exported", exported),
        ("", ""),
        ("Buildings in project", s["buildings"]),
        ("Assigned to a street", f"{s['assigned_to_street']} ({s['assigned_pct']}%)"),
        ("Unassigned", s["unassigned"]),
        ("Streets", s["streets"]),
        ("Streets named", s["streets_named"]),
        ("Streets awaiting a name", s["streets_unnamed"]),
        ("Total footprint m2", s["total_footprint_sqm"]),
        ("Median footprint m2", s["median_footprint_sqm"]),
        ("", ""),
        ("Buildings surveyed", s["buildings_surveyed"]),
        ("Premises estimated", s["premises_estimated_total"]),
        ("Premises basis", s["premises_basis"]),
    ]
    r = 4
    for label, value in rows:
        if label:
            ws.cell(row=r, column=1, value=label).font = Font(name="DM Sans", size=10,
                                                              color=STEEL)
            ws.cell(row=r, column=2, value=value).font = Font(name="DM Sans", size=10,
                                                              color=NAVY)
        r += 1

    r += 1
    _title(ws, "IMPORTANT", r)
    r += 1
    for note in [
        "Premises counts are estimates unless the row shows a surveyed unit "
        "count. They are not verified field data.",
        "Verification state indicates how strongly each record is evidenced. "
        "Records at 'imported' have not been checked.",
        m.as_text().splitlines()[-1] if m.blocked else
        "See the Sources sheet for attribution requirements.",
    ]:
        c = ws.cell(row=r, column=1, value=note)
        c.font = Font(name="DM Sans", size=9, color=STEEL)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 28
        r += 1


def _register_sheet(ws, rows: list[dict]) -> None:
    for i, (label, _, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=label)
        c.font = Font(name="Space Mono", size=9, bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"

    for r, row in enumerate(rows, start=2):
        for i, (_, key, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=i, value=row[key])
            cell.font = Font(name="DM Sans", size=9)
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="FFF7FAFD")
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"


def _sources_sheet(ws, m: Manifest) -> None:
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 14

    _title(ws, "DATA SOURCES AND ATTRIBUTION", 1)
    for i, label in enumerate(["Source", "Licence", "Licence class", "Features"],
                              start=1):
        c = ws.cell(row=3, column=i, value=label)
        c.font = Font(name="Space Mono", size=9, bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=BRAND)

    for r, s in enumerate(m.sources, start=4):
        ws.cell(row=r, column=1, value=s.name).font = Font(name="DM Sans", size=10)
        ws.cell(row=r, column=2,
                value=s.licence or "—").font = Font(name="DM Sans", size=10)
        ws.cell(row=r, column=3,
                value=s.licence_class.replace("_", " ")).font = Font(name="DM Sans", size=10)
        ws.cell(row=r, column=4, value=s.feature_count).font = Font(name="DM Sans", size=10)

    r = len(m.sources) + 6
    c = ws.cell(row=r, column=1, value=m.as_text())
    c.font = Font(name="DM Sans", size=9, color=STEEL)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 8, end_column=4)


def to_geojson(db: Session, user: User, project: Project, f: RegisterFilter,
               purpose: ExportPurpose) -> str:
    m = manifest(db, project, purpose)
    _guard(m)
    rows = _rows(db, project.id, f)
    features = []
    for r in rows:
        props = {k: v for k, v in r.items() if k != "_geom"}
        features.append({"type": "Feature",
                         "geometry": mapping(to_shape(r["_geom"])),
                         "properties": props})
    doc = {
        "type": "FeatureCollection",
        "features": features,
        "metadata": {
            "project": project.name,
            "generated": m.generated,
            "purpose": m.purpose.value,
            "filter": f.describe(),
            "crs_storage": "EPSG:4326",
            "crs_measurement": f"EPSG:{project.metric_crs_epsg}",
            "attribution": m.attribution_lines(),
        },
    }
    _audit(db, user, project, "geojson", len(rows), f, purpose)
    return json.dumps(doc)


def _audit(db: Session, user: User, project: Project, fmt: str, rows: int,
           f: RegisterFilter, purpose: ExportPurpose) -> None:
    audit_service.record(
        db, actor=user, entity_type="building", entity_id=None, action="export",
        project_id=project.id,
        changes={"export": {"before": None,
                            "after": {"format": fmt, "rows": rows,
                                      "purpose": purpose.value,
                                      "filter": f.describe()}}},
    )
    db.commit()


def walkthrough_pack(db: Session, user: User, project: Project) -> bytes:
    """The pack a surveyor carries: where to go, what to name, what to confirm.

    Internal use only — it is a work instruction, not a deliverable, so it
    carries no commercial licence gate.
    """
    currency = walkthrough_service.currency_report(db, project.id)
    grid = walkthrough_service.survey_grid(db, project)
    naming = walkthrough_service.naming_worklist(db, project.id)
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    wb = Workbook()
    ws = wb.active
    ws.title = "Brief"
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 46
    ws["A1"] = "AVIVA NETWORX"
    ws["A1"].font = Font(name="DM Sans", size=16, bold=True, color=NAVY)
    ws["A2"] = f"Field walkthrough pack — {project.district}"
    ws["A2"].font = Font(name="DM Sans", size=12, color=BRAND)

    brief = [
        ("Generated", generated),
        ("Project", project.name),
        ("Measurement CRS", f"EPSG:{project.metric_crs_epsg}"),
        ("", ""),
        ("Buildings in register", currency["total"]),
        ("Requiring field check", f'{currency["requires_field_check"]} '
                                  f'({currency["requires_field_check_pct"]}%)'),
        ("Median source age (years)", currency["median_age_years"]),
        ("Oldest source age (years)", currency["oldest_age_years"]),
        ("", ""),
        ("Priority 1 areas", f'{grid["priority_1_cells"]} cells, '
                             f'{grid["priority_1_buildings"]} buildings'),
        ("Priority 2 areas", f'{grid["priority_2_cells"]} cells, '
                             f'{grid["priority_2_buildings"]} buildings'),
        ("Streets awaiting a name", len(naming)),
    ]
    r = 4
    for label, value in brief:
        if label:
            ws.cell(row=r, column=1, value=label).font = Font(
                name="DM Sans", size=10, color=STEEL)
            ws.cell(row=r, column=2, value=value).font = Font(
                name="DM Sans", size=10, color=NAVY)
        r += 1

    r += 1
    _title(ws, "WHAT THIS PACK IS FOR", r)
    r += 1
    for note in [
        "Street names are deliberately blank. Record the name from the sign or "
        "from residents against the provisional code, and photograph the sign.",
        "The register was built from open data. Where it is wrong, it is wrong "
        "in a specific way: missing new build, and footprints altered since capture.",
        "Priority 1 areas rest on data over four years old or with no date at "
        "all. Walk these first and expect to add buildings, not just confirm them.",
        "Record the number of independently serviceable units per building. "
        "That figure is the one the whole capital plan depends on and no data "
        "source provides it.",
    ]:
        c = ws.cell(row=r, column=1, value=note)
        c.font = Font(name="DM Sans", size=9, color=STEEL)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.row_dimensions[r].height = 30
        r += 1

    _sheet(wb.create_sheet("1 — Streets to name"),
           ["Provisional code", "Road class", "Length m", "Buildings depending",
            "NAME OBSERVED", "Sign photographed?", "Surveyor", "Date"],
           [[n["provisional_code"], n["road_class"], round(n["length_m"]),
             n["buildings_depending"], "", "", "", ""] for n in naming],
           [18, 14, 10, 20, 32, 18, 18, 12])

    _sheet(wb.create_sheet("2 — Areas by priority"),
           ["Cell", "Priority", "Buildings", "Median age yr", "No date",
            "Unassigned", "Centre lat", "Centre lon", "Checked?"],
           [[c["cell"], c["survey_priority"], c["buildings"],
             c["median_age_years"], c["unknown_date"], c["unassigned"],
             c["centre_lat"], c["centre_lon"], ""] for c in grid["cells"]],
           [12, 10, 11, 14, 10, 12, 12, 12, 11])

    _sheet(wb.create_sheet("3 — Currency bands"),
           ["Band", "Buildings", "Survey priority", "Guidance"],
           [[currency["bands"][k]["label"], v,
             currency["bands"][k]["survey_priority"],
             currency["bands"][k]["note"]]
            for k, v in currency["by_currency"].items()],
           [16, 12, 16, 62])

    out = io.BytesIO()
    wb.save(out)
    audit_service.record(
        db, actor=user, entity_type="project", entity_id=project.id,
        action="export", project_id=project.id,
        changes={"export": {"before": None,
                            "after": {"format": "walkthrough_pack",
                                      "streets_to_name": len(naming),
                                      "priority_1_cells": grid["priority_1_cells"]}}},
    )
    db.commit()
    return out.getvalue()


def _sheet(ws, headers: list[str], rows: list[list], widths: list[int]) -> None:
    for i, (label, width) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=label)
        c.font = Font(name="Space Mono", size=9, bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"
    for r, row in enumerate(rows, start=2):
        for i, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=i, value=value)
            cell.font = Font(name="DM Sans", size=9)
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="FFF7FAFD")
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
