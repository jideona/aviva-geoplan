"""Excel export of the FAT coverage & density report."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "FF0D1B4B"; BRAND = "FF1A6FA8"; TEAL = "FF00C9A7"
AMBER = "FFB45309"; STEEL = "FF5A739A"; LIGHT = "FFF7FAFD"

TYPE_LABEL = {
    "single_home": "Single home", "secondary_structure": "Secondary / BQ",
    "duplex": "Duplex", "terrace": "Terrace", "apartment_block": "Apartment block",
    "commercial": "Commercial", "unknown": "Unknown",
}


def _hdr(ws, headers, widths, row=1, fill=NAVY):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name="Space Mono", size=9, bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 26


def build_workbook(report: dict) -> bytes:
    wb = Workbook()

    # --- summary ---
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 30
    ws["A1"] = "AVIVA NETWORX"
    ws["A1"].font = Font(name="DM Sans", size=16, bold=True, color=NAVY)
    ws["A2"] = "FAT coverage & property density report"
    ws["A2"].font = Font(name="DM Sans", size=12, color=BRAND)

    t = report["totals"]
    rows = [
        ("Project", report["project"]),
        ("Design engine", report["engine_version"]),
        ("FAT serving zones", report["fat_count"]),
        ("", ""),
        ("Buildings covered", t["buildings"]),
        ("  confirmed (surveyed)", t["confirmed_buildings"]),
        ("  estimated (inferred)", t["estimated_buildings"]),
        ("", ""),
        ("Confirmed units", t["confirmed_units"]),
        ("Estimated units (likely)", t["estimated_units_likely"]),
        ("Estimated units (range)",
         f'{t["estimated_units_low"]} – {t["estimated_units_high"]}'),
        ("Total units (likely)", t["total_units_likely"]),
    ]
    r = 4
    for label, value in rows:
        if label:
            ws.cell(row=r, column=1, value=label).font = Font(
                name="DM Sans", size=10, color=STEEL)
            ws.cell(row=r, column=2, value=value).font = Font(
                name="DM Sans", size=10, color=NAVY)
        r += 1
    r += 1
    note = ws.cell(row=r, column=1, value=report["note"])
    note.font = Font(name="DM Sans", size=9, italic=True, color=STEEL)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 3, end_column=2)

    # --- per-FAT ---
    fs = wb.create_sheet("FAT coverage")
    _hdr(fs, ["FAT", "Area m²", "Density /ha", "Quality", "Buildings",
              "Conf. bldg", "Est. bldg", "Conf. units", "Est. units",
              "Est. range", "Total likely", "Confidence", "Estates served"],
         [14, 10, 11, 10, 10, 10, 9, 11, 10, 12, 12, 11, 40])
    fs.freeze_panes = "A2"
    for i, f in enumerate(report["fats"], start=2):
        vals = [f["fat_code"], f["area_sqm"], f["density_per_ha"],
                f["coverage_quality"], f["buildings"], f["confirmed_buildings"],
                f["estimated_buildings"], f["confirmed_units"],
                f["estimated_units_likely"],
                f'{f["estimated_units_range"][0]}–{f["estimated_units_range"][1]}',
                f["total_units_likely"], f["overall_confidence"],
                ", ".join(f["estates_served"])]
        for j, v in enumerate(vals, start=1):
            c = fs.cell(row=i, column=j, value=v)
            c.font = Font(name="DM Sans", size=9)
            if i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=LIGHT)
        # colour the quality cell
        q = f["coverage_quality"]
        qc = fs.cell(row=i, column=4)
        qc.font = Font(name="Space Mono", size=9, bold=True,
                       color=TEAL if q == "confirmed" else
                       AMBER if q == "estimated" else BRAND)
    fs.auto_filter.ref = f"A1:M{len(report['fats']) + 1}"

    # --- type breakdown, one row per (FAT, type) ---
    bs = wb.create_sheet("Type breakdown")
    _hdr(bs, ["FAT", "Property type", "Buildings", "Confirmed", "Estimated",
              "Units likely", "Units range", "Confidence", "Data quality"],
         [14, 18, 10, 10, 10, 12, 12, 11, 14], fill=BRAND)
    bs.freeze_panes = "A2"
    row = 2
    for f in report["fats"]:
        for b in f["breakdown"]:
            quality = ("confirmed" if b["confirmed_buildings"] == b["buildings"]
                       else "estimated" if b["confirmed_buildings"] == 0
                       else "mixed")
            vals = [f["fat_code"], TYPE_LABEL.get(b["type"], b["type"]),
                    b["buildings"], b["confirmed_buildings"],
                    b["estimated_buildings"], b["units_likely"],
                    f'{b["units_range"][0]}–{b["units_range"][1]}',
                    b["confidence"], quality]
            for j, v in enumerate(vals, start=1):
                c = bs.cell(row=row, column=j, value=v)
                c.font = Font(name="DM Sans", size=9)
                if quality == "estimated" and j == 9:
                    c.font = Font(name="Space Mono", size=9, color=AMBER)
            row += 1
    bs.auto_filter.ref = f"A1:I{row - 1}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
